"""Generate Kings Equestrian Inventory dashboard user guide (.docx)."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

OUT = os.path.join(os.path.dirname(__file__), "Kings_Equestrian_Inventory_User_Guide.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def title(text):
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def h1(text):
    doc.add_heading(text, level=1)

def h2(text):
    doc.add_heading(text, level=2)

def para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")

title("Kings Equestrian Inventory System")
para("Dashboard User Guide — activities by tab", bold=True)
para("Version: vendor procurement workflow (current app)")
doc.add_paragraph()

h1("1. About this app")
para(
    "This web dashboard manages inventory procurement and stock for Kings Equestrian "
    "(locations: Bangalore, Hyderabad, Pune, Farm). Parent or customer orders are handled "
    "on other platforms; this app tracks buying from vendors, receiving stock, and internal "
    "inventory operations."
)
para("Sign in with USER_MASTER credentials (username/email + password) — not your Google account.")

h1("2. User roles & tab visibility")
bullets([
    "Admin — full access to all tabs and all locations; can switch View location in the header.",
    "Accounts — Summary view; Pay vendor tab; can record vendor payments and auto-approve requests.",
    "Trainer — Summary view; Request tab; can create procurement requests and manage samples.",
    "Manager — location-scoped access to masters (vendor/item), orders, receive, issue, transfer.",
])
para("Non-admin users see only data for their assigned location. Location fields in forms are locked for them.")

h1("3. End-to-end procurement workflow")
numbered([
    "Request — raise a procurement need (vendor + item + quantity + location). Example: 20 Helmets from Saif at Hyderabad.",
    "Pay vendor — Accounts records payment to the vendor; request status becomes Approved.",
    "Order — Manager/Admin places the vendor purchase order (PO) for the location.",
    "Receive — goods are received from the vendor; inventory quantity increases.",
    "Issue (optional) — stock is issued internally after goods have been received.",
])
para("Business rules enforced by the system:", bold=True)
bullets([
    "No vendor payment → request cannot be approved for ordering.",
    "No approval → vendor order cannot be placed.",
    "No goods receipt → item cannot be issued against that request.",
    "Issue is blocked if available stock is less than issue quantity.",
])

h1("4. Activities by tab")

tabs = [
    ("Summary", "All signed-in users",
     "Dashboard home with KPIs and stock overview.",
     [
         "View KPI cards: total SKUs, pending requests, active requests, vendor payments, orders, low-stock count.",
         "Click a KPI card to jump to the related tab.",
         "Fulfilment pipeline: counts for Pending → Approved → Ordered → Received → Issued.",
         "Inventory table: all SKUs for the selected View location (item code, name, qty, min level, location).",
         "Low stock alerts: items at or below minimum level.",
         "Refresh data — reload summary from the spreadsheet.",
     ],
     "Admin can change View (All Locations or a specific site) in the header. Others see their home location only."),
    ("Inventory", "All signed-in users",
     "Advanced inventory analytics and storage setup.",
     [
         "Inventory Control Center — full stock list with totals for the selected location.",
         "Location analytics — stock breakdown by site.",
         "Inventory ownership — breakdown by inventory type (Business, Procurement, Sample/Demo, Operational).",
         "Fast moving — items with recent movement.",
         "Dead stock — items with no movement for 90+ days.",
         "Add storage location — register shelves/rooms (e.g. HYD Shelf 1) tied to a location.",
     ],
     None),
    ("Vendor ➕", "Admin, Manager (master permission)",
     "Add suppliers to VENDOR_MASTER.",
     [
         "Enter vendor name, short code (3–5 letters, used in item codes), phone, email, location.",
         "Save vendor — creates a new vendor row (e.g. VEN-0001).",
     ],
     "Vendor code is used when auto-generating item codes: KE-[Code]-[Category]-[Model]."),
    ("New Item ➕", "Admin, Manager (master permission)",
     "Add SKUs to INVENTORY_MASTER.",
     [
         "Item name, category, model, vendor, location.",
         "Opening quantity and minimum stock level.",
         "Inventory type and optional allocated-to field.",
         "Storage location (from registered storage names).",
         "Cost price — Admin only.",
         "Image — URL or upload from camera/gallery (stored on Google Drive).",
         "Save item — auto-generates item code and inventory row.",
     ],
     None),
    ("Request", "Admin, Trainer",
     "Procurement requests to vendors (REQUEST_REGISTER).",
     [
         "View procurement requests register — filter by date range and status; paginated table.",
         "Create procurement request — select Vendor, Item, Quantity, Location.",
         "Example: Saif · Helmet · 20 · Hyderabad.",
         "Creates REQ-* row with status Pending and Payment Status Unpaid.",
     ],
     "Items must already exist in INVENTORY_MASTER for the chosen location."),
    ("Pay vendor", "Admin, Accounts",
     "Record payments made to vendors (PAYMENTS).",
     [
         "View vendor payments register — filter by date and status.",
         "Pay vendor form — select Request ID (auto-fills vendor from request).",
         "Enter amount and payment mode (Bank, UPI, Card, Cash).",
         "Record vendor payment & approve — writes PAY-* row and sets request to Approved.",
     ],
     "Required before a vendor order can be placed."),
    ("Order", "Admin, Manager (location-scoped)",
     "Place vendor purchase orders (VENDOR_ORDERS).",
     [
         "View vendor orders register — filter by date and status.",
         "Place vendor order — select Request ID (auto-fills vendor, item, qty, location).",
         "Confirm vendor, item, quantity, location.",
         "Place order — creates ORD-* row; request status becomes Ordered.",
     ],
     "Request must be Approved (vendor payment recorded)."),
    ("Receive", "Admin, Manager (location-scoped)",
     "Goods received from vendor (GOODS_RECEIVED / GRN).",
     [
         "View goods received register — filter by date and status.",
         "Receive goods — select Order ID, quantity received, location, storage location.",
         "Receive & update stock — creates GRN-* row and increases INVENTORY_MASTER quantity.",
         "Updates linked request status to Received.",
     ],
     "This is how stock enters the system after a vendor delivery."),
    ("Issue", "Admin, Manager (location-scoped)",
     "Issue stock internally (ISSUE_REGISTER).",
     [
         "View issues register — filter by date and status.",
         "Issue item — select Request ID, issued to (name), quantity, location.",
         "Issue item — reduces inventory; completes the request workflow.",
     ],
     "Goods must be received against the request’s order before issue is allowed."),
    ("Transfer", "Admin, Manager",
     "Move stock between locations (STOCK_TRANSFER).",
     [
         "Move stock — select item, from location, to location, quantity, optional notes.",
         "Move stock now — creates transfer record and adjusts stock at both sites.",
         "Transfer log — filter by date, status, item, from/to locations.",
     ],
     None),
    ("Sample", "Admin, Manager, Trainer",
     "Sample / demo inventory (SAMPLE_MOVEMENT).",
     [
         "Issue sample — item, location, qty, issued to, purpose, expected return date, notes.",
         "Issue sample & reduce stock — tracks demo units leaving stock.",
         "Sample register — filter and review Out / Returned / Lost samples.",
         "Return or mark lost actions available from the register (Manager/Admin).",
     ],
     None),
    ("Pricing", "Admin only",
     "Selling prices and margins (ITEM_PRICING_MASTER).",
     [
         "Set cost price, selling price, MRP, preferred vendor for an item.",
         "Save pricing — updates pricing master.",
         "View pricing master table with filters.",
     ],
     "Operations staff do not see cost or margin fields."),
    ("Finance", "Admin only",
     "Financial KPIs and reporting.",
     [
         "Financial KPIs — revenue/payables-style summary for the selected location.",
         "Vendor profitability table.",
         "Send weekly report to admins now (on-demand email).",
     ],
     "Weekly auto-email requires installWeeklyReportTrigger() in Apps Script."),
]

for name, who, purpose, activities, note in tabs:
    h2(f"Tab: {name}")
    para(f"Who can access: {who}")
    para(f"Purpose: {purpose}")
    para("Activities you can do:", bold=True)
    bullets(activities)
    if note:
        para(f"Note: {note}")
    doc.add_paragraph()

h1("5. Registers (tables in Request / Payment / Order / Receive / Issue tabs)")
para("Each procurement tab includes a register at the top with common controls:")
bullets([
    "Date from / Date to — filter rows by date.",
    "Status dropdown — filter by workflow status.",
    "Apply — run filter; Clear — reset filters.",
    "Pagination — browse older records.",
])

h1("6. Header controls")
bullets([
    "View — Admin selects All Locations or a specific site; filters all data.",
    "Role pill — shows your role (Admin, Accounts, Trainer, Manager).",
    "User chip — your display name.",
    "Sign out — ends session.",
])

h1("7. Related Google Sheets (backend)")
para("Data is stored in the linked spreadsheet. Key tabs:")
bullets([
    "INVENTORY_MASTER — stock SKUs and quantities.",
    "VENDOR_MASTER — suppliers.",
    "REQUEST_REGISTER — procurement requests.",
    "PAYMENTS — vendor payments.",
    "VENDOR_ORDERS — purchase orders.",
    "GOODS_RECEIVED — GRN / receipts.",
    "ISSUE_REGISTER — internal issues.",
    "USER_MASTER — login and roles.",
    "STOCK_TRANSFER, SAMPLE_MOVEMENT, ITEM_PRICING_MASTER, STORAGE_LOCATIONS — advanced modules.",
])

h1("8. Quick reference — example: order 20 helmets for Hyderabad")
numbered([
    "Vendor ➕ — ensure Saif exists in VENDOR_MASTER.",
    "New Item ➕ — ensure Helmet exists at Hyderabad in INVENTORY_MASTER.",
    "Request — Saif, Helmet, 20, Hyderabad.",
    "Pay vendor — select REQ-*, enter amount, record payment.",
    "Order — select same REQ-*, place vendor order.",
    "Receive — select ORD-*, enter qty 20, receive goods → stock +20.",
    "Summary / Inventory — verify Helmet qty at Hyderabad.",
])

doc.save(OUT)
print("Wrote", OUT)
